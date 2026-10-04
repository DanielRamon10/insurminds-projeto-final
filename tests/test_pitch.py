"""Testes do gerador do pitch deck (E.3).

O deck é entregável avaliado e o nome do arquivo é literal, então há dois tipos
de teste aqui:

* **o arquivo** — gera de verdade num diretório temporário e confere que o
  `.pptx` abre, tem os slides, o formato 16:9 e nenhum slide vazio;
* **o conteúdo** — confere que os números escritos no deck batem com o que o
  código produz hoje. Sem isso, o corretor acrescenta um campo ao dicionário e o
  deck continua dizendo "15 campos" na frente da banca.
"""

from __future__ import annotations

from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402

from app.domain.campos import carregar_campos  # noqa: E402
from app.domain.comparacao import Veredito, comparar  # noqa: E402
from app.domain.exemplos import carregar_exemplos  # noqa: E402
from app.interface.apresentacao import linhas_rastreabilidade  # noqa: E402
from scripts.gerar_pitch import CONTEUDO, DESTINO, montar  # noqa: E402


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    """O deck gerado uma vez, num diretório temporário."""
    caminho = tmp_path_factory.mktemp("pitch") / "InsurMinds_Projeto_Final.pptx"
    return montar(caminho)


def texto_do_deck(caminho: Path) -> str:
    apresentacao = Presentation(str(caminho))
    return "\n".join(
        forma.text_frame.text
        for slide in apresentacao.slides
        for forma in slide.shapes
        if forma.has_text_frame
    )


def slide_por_tipo(tipo: str) -> dict:
    return next(s for s in CONTEUDO if s["tipo"] == tipo)


# ---------------------------------------------------------------------------
# O arquivo
# ---------------------------------------------------------------------------


def test_nome_do_arquivo_e_o_exigido_pelo_enunciado():
    assert DESTINO.name == "InsurMinds_Projeto_Final.pptx"
    assert DESTINO.parent.name == "Projeto_Final_Artefatos"


def test_gera_um_pptx_valido_com_todos_os_slides(deck):
    apresentacao = Presentation(str(deck))
    assert len(apresentacao.slides) == len(CONTEUDO)
    assert len(apresentacao.slides) >= 12


def test_formato_de_apresentacao_e_16_9(deck):
    apresentacao = Presentation(str(deck))
    assert round(apresentacao.slide_width / apresentacao.slide_height, 2) == 1.78


def test_nenhum_slide_saiu_vazio(deck):
    """Um slide em branco no meio do deck só se descobre na hora de apresentar."""
    apresentacao = Presentation(str(deck))
    for numero, slide in enumerate(apresentacao.slides, start=1):
        textos = [f.text_frame.text for f in slide.shapes if f.has_text_frame]
        assert any(texto.strip() for texto in textos), f"slide {numero} saiu sem texto"


# ---------------------------------------------------------------------------
# O conteúdo
# ---------------------------------------------------------------------------


def test_o_deck_nomeia_todos_os_integrantes(deck):
    texto = texto_do_deck(deck)
    for nome in (
        "Daniel Ramon",
        "Paulo Henrique",
        "Nicole Paes",
        "Paulo Roberto",
        "Juliana Catarina",
    ):
        assert nome in texto, f"{nome} nao aparece no deck"


def test_a_capa_declara_a_entrega(deck):
    capa = Presentation(str(deck)).slides[0]
    texto = "\n".join(f.text_frame.text for f in capa.shapes if f.has_text_frame)
    assert "06/10/2026" in texto
    # Jarvis e o grupo; InsurMinds e o curso. A capa traz os dois, e confundi-los
    # foi exatamente o erro que este teste passou a cobrir. A comparacao ignora a
    # caixa porque o kicker da capa e escrito em maiusculas.
    assert "jarvis" in texto.lower()
    assert "insurminds" in texto.lower()


def test_o_deck_declara_a_limitacao_real_da_extracao(deck):
    """O deck não pode vender como pronto o que não está.

    A versão anterior deste teste exigia a palavra "desenvolvimento" no deck,
    porque a frente B ainda não estava publicada. Com `extrair_apolice` em
    `app.agents`, essa frase passou a ser falsa — e um teste que exige uma frase
    falsa é pior do que nenhum teste, porque trava a correção. O que o deck tem
    de declarar agora é a limitação que existe de verdade: os campos numéricos
    não têm valor nas condições gerais — elas os remetem à Especificação da
    Apólice —, e é por isso que a demonstração roda sobre as extrações de exemplo.
    """
    texto = texto_do_deck(deck)
    assert "Especificação da Apólice" in texto, (
        "o deck não declara por que os campos numéricos não têm valor: as condições "
        "gerais remetem cada um à Especificação da Apólice"
    )
    assert "exemplo" in texto, (
        "o deck não declara que a demonstração roda sobre extrações de exemplo"
    )
    assert "em desenvolvimento" not in texto, (
        "a extração está publicada em `app.agents` (frente B): o deck não pode "
        "dizer que ela está em desenvolvimento"
    )


def test_o_numero_de_campos_do_deck_bate_com_o_dicionario():
    dicionario = carregar_campos()
    slide = next(s for s in CONTEUDO if s.get("titulo", "").startswith("O que comparar"))
    valores = [numero[0] for numero in slide["numeros"]]
    assert str(len(dicionario)) in valores, (
        f"o deck diz {valores} e o dicionario tem {len(dicionario)} campos"
    )


def test_os_numeros_do_slide_de_resultados_batem_com_o_motor():
    """O slide de resultados não pode envelhecer em relação ao código.

    A fonte aqui é `exemplos.py`, de propósito: é o que a tela mostra quando o
    banco só tem as extrações de exemplo, que é o caso da gravação descrita em
    `docs/ROTEIRO_VIDEO.md`. A extração real de `data/extracoes/` devolve 9 e 10
    dos 15 campos — e nenhum dos cinco numéricos —, então trocar a fonte sem
    decidir isso antes deixaria o deck dizendo números que a tela não mostra.
    """
    apolices = carregar_exemplos()
    dicionario = carregar_campos()
    resultado = comparar(apolices, dicionario)
    linhas = linhas_rastreabilidade(resultado, apolices)

    slide = next(s for s in CONTEUDO if s.get("titulo") == "O resultado, campo a campo")
    valores = [numero[0] for numero in slide["numeros"]]

    for veredito in (
        Veredito.AUSENTE_EM_ALGUMA,
        Veredito.DIFERENTE,
        Veredito.REDACAO_DIVERGENTE,
    ):
        assert str(resultado.resumo[veredito.value]) in valores, (
            f"o deck nao cita o total de '{veredito.value}'"
        )
    assert str(len(linhas)) in valores


def test_todo_slide_de_conteudo_tem_titulo_e_kicker():
    for dados in CONTEUDO:
        if dados["tipo"] == "capa":
            continue
        assert dados["kicker"].strip()
        assert dados["titulo"].strip()


def test_o_slide_de_arquitetura_cobre_as_quatro_frentes():
    slide = slide_por_tipo("fluxo")
    siglas = [etapa[0] for etapa in slide["etapas"]]
    assert siglas == ["A", "B", "C", "D"]


def test_o_slide_de_limites_avisa_do_risco_do_projeto():
    """A limitação mais perigosa é a extração inventar valor."""
    slide = next(s for s in CONTEUDO if s.get("kicker") == "limitações conhecidas")
    texto = " ".join(titulo + corpo for _, titulo, corpo in slide["topicos"])
    assert "LLM" in texto or "modelo" in texto


# ---------------------------------------------------------------------------
# O slide de demonstração — os recortes da interface
# ---------------------------------------------------------------------------


def test_ha_um_print_por_legenda():
    """Um arquivo por legenda, na mesma ordem.

    Se as duas listas desalinharem, o slide cola a tela da tabela embaixo do
    texto que fala da comparação — e ninguém percebe olhando o deck, porque as
    três telas são parecidas.
    """
    from scripts.gerar_pitch import ARQUIVOS_DE_PRINT

    assert len(ARQUIVOS_DE_PRINT) == len(slide_por_tipo("prints")["legendas"])


def test_o_slide_de_demonstracao_ou_cola_os_prints_ou_mostra_a_caixa_vazia(deck):
    """Nunca as duas coisas no mesmo slide.

    O slide é honesto nos dois estados: com os recortes feitos ele mostra as
    telas, e num clone recém-baixado ele mostra a caixa com "print N — colar
    aqui". O que ele não pode é misturar — três caixas de alturas diferentes,
    ou uma tela colada ao lado de duas caixas vazias, denuncia que a lista de
    arquivos e a de legendas saíram de sincronia.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    from scripts.gerar_pitch import ARQUIVOS_DE_PRINT, PRINTS

    indice = next(i for i, s in enumerate(CONTEUDO) if s["tipo"] == "prints")
    slide = Presentation(str(deck)).slides[indice]

    fotos = sum(1 for f in slide.shapes if f.shape_type == MSO_SHAPE_TYPE.PICTURE)
    colados = sum(1 for nome in ARQUIVOS_DE_PRINT if (PRINTS / nome).exists())

    assert fotos == colados, (
        f"{colados} print(s) em disco e {fotos} colado(s) no slide — "
        "o slide misturou os dois estados"
    )


def test_o_print_colado_nao_invade_a_legenda(deck):
    """A imagem cabe entre a chamada do slide e o título da legenda.

    A legenda começa em 4,98". Se a imagem passar disso ela cobre o próprio
    texto que a explica — e o defeito só aparece abrindo o deck.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    indice = next(i for i, s in enumerate(CONTEUDO) if s["tipo"] == "prints")
    slide = Presentation(str(deck)).slides[indice]

    for foto in (f for f in slide.shapes if f.shape_type == MSO_SHAPE_TYPE.PICTURE):
        base = (foto.top + foto.height) / 914400
        assert base <= 4.98, f"a imagem termina em {base:.2f}\" e a legenda começa em 4,98\""
        assert foto.top / 914400 >= 2.25, "a imagem invade a chamada do slide"
